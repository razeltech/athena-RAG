from app.core.parser import DocumentParser, ParsedDocument


class XlsxParser(DocumentParser):
    @property
    def extensions(self) -> list[str]:
        return [".xlsx"]

    def extract(self, path: str, source: str) -> ParsedDocument:
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        lines: list[str] = []
        for sheet in wb.worksheets:
            lines.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if cells:
                    lines.append(" | ".join(cells))
        return ParsedDocument(text="\n".join(lines), source=source)
